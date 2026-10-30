"""
Fleet audit: how Energex distribution switches match IPS setting records.

Offline - no PowerFactory. Reads a description of every switch, either

    * one CSV from export_pf_switches.py, or
    * a folder of the per-project match logs IPStoPF writes during a batch
      (config/switch_matching.py, folder "switch_match"),

matches them with mapping/switch_matcher.py against the IPS setting-ID
cache, compares with the legacy name lookup, and writes a report folder:

    summary.md                   headline counts and what to look at
    switch_match_audit.xlsx      all the tables below, one sheet each
    switches.csv                 every switch: legacy result, new result, rule
    relays_at_risk.csv           processed switches with relays/fuses in the
                                 cubicle but no IPS match - IPStoPF deletes
                                 those relays
    ips_unmatched.csv            in-scope IPS records no switch takes, with a
                                 reason
    filtered_with_ips.csv        switches IPStoPF skips (de-energised, out of
                                 service, name rule) that IPS has settings for
    review.csv                   ambiguous matches, rejected candidates,
                                 override problems
    legacy_changes.csv           switches whose records change vs. legacy
    suggested_overrides.csv      candidate rows for switch_match_overrides.csv

Usage (from the IPStoPF folder, so the package imports resolve):

    python switch_match_audit.py --switches "<export.csv or match-log folder>"
        [--ips Report-Cache-ProtectionSettingIDs-EX.csv]
        [--overrides <switch_match_overrides.csv>]
        [--out <report folder>]
        [--enable transformer_cb] [--disable composite]

Without --ips the cache is fetched the way the batch fetches it
(ips_data.query_database). Without --overrides the configured overrides
file is used (config/switch_matching.py).
"""

from __future__ import annotations

import argparse
import csv
import difflib
import os
import re
import sys
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Set, Tuple

sys.path.insert(0, str(Path(__file__).resolve().parent))

from config import switch_matching as smc  # noqa: E402
from mapping import switch_matcher as sm  # noqa: E402
from process_ips.ips_ingest import classify_designation  # noqa: E402
from process_ips.ips_records import IpsElementType  # noqa: E402

ALL_GRIDS_STUDY_CASE = "All Active Grids Study Case"

# Down-line switch name suffixes that denote a protection device.
PROTECTIVE_SUFFIXES = {"RE", "SE", "SW", "DO", "PF", "MDO", "FUSE"}

# Substation element types the distribution models do not carry as switches.
NOT_SWITCH_TYPES = {
    IpsElementType.BUSBAR,
    IpsElementType.CAPACITOR_BANK,
    IpsElementType.TRANSFORMER,
    IpsElementType.TRANSFORMER_SWITCH,
    IpsElementType.BUS_COUPLER,
    IpsElementType.FEEDER,
}


# =============================================================================
# Loading
# =============================================================================

def _read_csv(path: Path) -> List[Dict[str, str]]:
    with open(path, newline="", encoding="utf-8-sig") as f:
        return [{k: (v or "") for k, v in row.items() if k} for row in csv.DictReader(f)]


def load_switch_rows(source: str) -> Tuple[List[Dict[str, str]], List[str]]:
    """Rows from an export CSV, or from every switch_match_*.csv in a folder."""
    path = Path(source)
    if path.is_dir():
        files = sorted(path.glob("switch_match_*.csv"))
        if not files:
            raise SystemExit(f"No switch_match_*.csv files in {path}")
    else:
        files = [path]
    rows: List[Dict[str, str]] = []
    for f in files:
        rows.extend(_read_csv(f))
    return rows, [str(f) for f in files]


def load_ips_rows(path: Optional[str]) -> Tuple[List[Dict[str, Any]], str]:
    if path:
        return _read_csv(Path(path)), path
    from ips_data.query_database import _create_ids_dict  # needs the network
    rows = _create_ids_dict("Energex")
    if not rows:
        raise SystemExit("Could not fetch the IPS cache; pass --ips <csv>")
    return rows, "corporate cache (Report-Cache-ProtectionSettingIDs-EX)"


# NEW_NAME values that are review notes, not names (ips_data.cb_mapping
# skips the first four; the others find nothing in IPS).
CB_ALT_NOTES = {"not needed", "no active setting", "wrong sub name", "unknown",
                "not required", "decommisioned", "decommissioned"}


def load_cb_alt_rows(path: Optional[str] = None) -> List[Dict[str, str]]:
    """Every row of CB_ALT_NAME.csv, review notes included, with its line."""
    if path is None:
        from config.paths import get_cb_alt_name_file
        path = str(get_cb_alt_name_file())
    try:
        rows = _read_csv(Path(path))
    except FileNotFoundError:
        return []
    for n, row in enumerate(rows, start=2):
        row["line"] = str(n)
    return rows


def cb_alt_index(rows: List[Dict[str, str]]) -> Dict[Tuple[str, str], str]:
    """The CB_ALT_NAME lookup the matcher uses: notes skipped, first row per key wins."""
    index: Dict[Tuple[str, str], str] = {}
    for row in rows:
        if row.get("NEW_NAME", "").strip().lower() in CB_ALT_NOTES - {
                "not required", "decommisioned", "decommissioned"}:
            continue
        index.setdefault((row.get("SUBSTATION"), row.get("CB_NAME")), row.get("NEW_NAME"))
    return index


def _kv(value: str) -> Optional[float]:
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def to_ref(row: Dict[str, str]) -> sm.SwitchRef:
    """SwitchRef from an export / match-log row (as ex_switch_matching builds it)."""
    if row.get("class") == "ElmCoup":
        container, container_class, substation = (
            row.get("fold_id_name", ""), row.get("fold_id_class", ""), None
        )
    else:
        container, container_class = row.get("terminal_name", ""), "ElmTerm"
        substation = row.get("substation") or None
    return sm.SwitchRef(
        key=row.get("full_name", ""),
        project=row.get("project", ""),
        cls=row.get("class", ""),
        name=row.get("loc_name", ""),
        container=container,
        container_class=container_class,
        fold_name=row.get("fold_id_name", ""),
        kv=_kv(row.get("terminal_kv", "")),
        substation=substation,
    )


def is_processed(row: Dict[str, str]) -> bool:
    return row.get("valid_switch") == "1" and row.get("should_process") == "1"


def filter_reason(row: Dict[str, str]) -> str:
    """Why IPStoPF skips a switch (mirrors _get_valid_switches / _should_process_switch)."""
    reasons = []
    if row.get("valid_switch") != "1":
        if not row.get("grid"):
            reasons.append("not in a grid")
        if row.get("out_of_service") == "1":
            reasons.append("out of service")
        if row.get("calc_relevant") == "0":
            case = row.get("study_case", "")
            if case and case != ALL_GRIDS_STUDY_CASE:
                reasons.append(f"grid not active in study case '{case}'")
            else:
                reasons.append("not calculation relevant")
        if row.get("energised") == "0":
            reasons.append("de-energised")
        if not reasons:
            reasons.append("not energised / not relevant in the active study case")
    if row.get("should_process") != "1":
        if row.get("class") == "ElmCoup":
            reasons.append("ElmCoup name has a digit as 3rd character (RMU switch rule)")
        else:
            reasons.append("cubicle not connected to an element")
    return "; ".join(reasons)


# =============================================================================
# Legacy lookup (the real SettingIndex code)
# =============================================================================

def legacy_ids(setting_index, row: Dict[str, str], cb_alt: Dict) -> List[str]:
    """
    The old Energex name lookup (removed from ex_settings on 2026-10-30:
    _get_switch_info + _get_setting_id_indexed), kept here for the
    Changes vs legacy comparison.
    """
    new_name = cb_alt.get((row.get("fold_id_name"), row.get("loc_name")))
    name = (new_name if new_name is not None else row.get("loc_name", "")).split("_")[0]
    sub_code = row.get("fold_id_name") if row.get("class") == "ElmCoup" else None
    if len(name) < 4:
        return []
    return [r.relaysettingid for r in setting_index.get_by_switch_name(name, sub_code)]


# =============================================================================
# Suggestions
# =============================================================================

_SITE_CORE = re.compile(r"[^A-Z0-9]")


def _device_names(row: Dict[str, str]) -> List[str]:
    out = []
    for col in ("relays", "fuses"):
        for item in row.get(col, "").split(";"):
            item = item.strip()
            if item:
                out.append(re.sub(r"\s*\[[^\]]*\]$", "", item))
    return out


def _trailing_number(name: str) -> Optional[str]:
    m = re.search(r"(\d+)[A-Z]?$", sm._DUP_SUFFIX.sub("", name.split("_")[0]))
    return m.group(1) if m else None


class _Lookup:
    """IPS entries by upper-cased name and by site, for suggestions."""

    def __init__(self, entries: Iterable[sm.IpsEntry]):
        self.by_name: Dict[str, List[sm.IpsEntry]] = defaultdict(list)
        self.names_at: Dict[str, Set[str]] = defaultdict(set)
        for e in entries:
            name = e.name.strip().upper()
            self.by_name[name].append(e)
            if e.site:
                self.names_at[e.site].add(name)
        self.squeezed_dld: Dict[str, Set[str]] = defaultdict(set)
        for e in entries:
            if e.category == sm.CATEGORY_DOWN_LINE:
                name = e.name.strip().upper()
                self.squeezed_dld[_SITE_CORE.sub("", name)].add(name)


_TRANSFORMER_TAIL = re.compile(r"^\d[XT]\d+$|^TR?\d+$")


def suggest(
    row: Dict[str, str],
    ref: sm.SwitchRef,
    res: sm.SwitchMatch,
    lookup: _Lookup,
    claimed: Set[str],
    off_rules: Set[str],
    off_matcher: Optional[sm.SwitchMatcher],
) -> List[Dict[str, str]]:
    """
    Candidate override rows for one unmatched or ambiguous switch.

    Never applied: these go to suggested_overrides.csv for review.
    Confidence:
        high   - the cubicle already holds relays named after the IPS device
        medium - one similar name at the same site, a rule that is switched
                 off would match, or a Down Line Device written differently
        low    - several candidates; pick one (or none)
    """
    found: List[Tuple[str, str, str, List[sm.IpsEntry]]] = []
    by_name = lookup.by_name

    def located(entries):
        return [e for e in entries
                if sm.location_problem(ref, e) is None and e.setting_id not in claimed]

    # 0. Ambiguous: the rule found the record, but another switch competes.
    if res.status == sm.AMBIGUOUS and res.rejections:
        rej = res.rejections[-1]
        entries = [e for e in located(by_name.get(_norm(rej.pf_name), []))]
        if entries:
            found.append(("low", rej.pf_name,
                          f"{res.note}; pick the switch that carries it", entries))

    # 1. The cubicle already holds relays named after an IPS device.
    for dev in _device_names(row):
        parts = dev.split("_")
        for k in range(len(parts) - 1, 0, -1):
            cand = "_".join(parts[:k]).strip().upper()
            entries = located(by_name.get(cand, []))
            if entries:
                found.append(("high", cand,
                              f"cubicle already holds '{dev}', named for this IPS device",
                              entries))
                break

    core = sm._DUP_SUFFIX.sub("", ref.name.split("_")[0].strip()).upper()
    tail = core[len(ref.container):] if core.startswith(ref.container) else ""
    is_transformer = ref.cls == "ElmCoup" and bool(_TRANSFORMER_TAIL.match(tail))

    # 2. A rule that is switched off (e.g. transformer_cb) would match.
    if not found and off_matcher is not None:
        alt = off_matcher.propose(ref)
        if alt.status == sm.MATCHED and alt.pf_rule in off_rules:
            ok = located(alt.entries)
            if ok:
                found.append(("medium", alt.pf_name,
                              f"rule '{alt.pf_rule}' (switched off) would match",
                              ok))

    # 3. One similar name at the same site.
    if not found and ref.cls == "ElmCoup" and len(core) >= 4 and not is_transformer:
        num = _trailing_number(core)
        scored = []
        for name in sorted(lookup.names_at.get(ref.container, ())):
            ok = located(by_name[name])
            if not ok:
                continue
            tokens = [name] + [t.upper() for t in sm._composite_keys(name)]
            score = max(difflib.SequenceMatcher(None, core, t).ratio() for t in tokens)
            if score >= 0.75 and (num is None or any(_trailing_number(t) == num for t in tokens)):
                scored.append((score, name, ok))
        scored.sort(key=lambda x: (-x[0], x[1]))
        for score, name, ok in scored[:3]:
            conf = "medium" if len(scored) == 1 and score >= 0.8 else "low"
            found.append((conf, name,
                          f"name similar to '{core}' at {ref.container} (score {score:.2f})"
                          + ("" if len(scored) == 1 else f"; {len(scored)} candidates"),
                          ok))

    # 4. A Down Line Device written differently (case, spaces, hyphens only;
    #    a different digit is a different pole).
    if not found and ref.cls == "StaSwitch":
        m = sm._DOWN_LINE_SLASH.match(ref.name.strip())
        base = (m.group(1) if m else ref.name.split("_")[0]).strip().upper()
        squeezed = _SITE_CORE.sub("", base)
        if len(squeezed) >= 5:
            for name in lookup.squeezed_dld.get(squeezed, ()):
                if name == base:
                    continue
                ok = located(by_name.get(name, []))
                if ok:
                    found.append(("medium", name,
                                  f"Down Line Device '{name}' is written '{base}' in PF", ok))

    # One row per (switch, IPS name, site).
    merged: Dict[Tuple[str, str], Dict[str, str]] = {}
    for conf, name, evidence, entries in found:
        key = (entries[0].name.strip(), entries[0].site or "")
        if key in merged:
            if evidence not in merged[key]["evidence"]:
                merged[key]["evidence"] += "; " + evidence
            continue
        merged[key] = {
            "project": ref.project,
            "switch_class": ref.cls,
            "container": ref.container,
            "switch_name": ref.name,
            "action": "MATCH",
            "ips_name": key[0],
            "ips_site": key[1],
            "note": "",
            "confidence": conf,
            "evidence": evidence,
            "ips_patterns": "; ".join(sorted({e.pattern for e in entries})),
            "switch_relays": row.get("relays", "") or row.get("fuses", ""),
        }
    return list(merged.values())


def _norm(name: str) -> str:
    return name.strip().upper()


# =============================================================================
# Audit
# =============================================================================

def _ips_label(e: sm.IpsEntry) -> str:
    return f"{e.name.strip()} @ {e.site} [{e.pattern}]"


def run_audit(
    switch_rows: List[Dict[str, str]],
    ips_rows: List[Dict[str, Any]],
    overrides: Sequence[sm.Override],
    cb_alt: Dict,
    rules: Set[str],
    cb_alt_rows: Optional[List[Dict[str, str]]] = None,
) -> Dict[str, Any]:
    from ips_data.setting_index import SettingIndex

    setting_index = SettingIndex(ips_rows, "Energex")
    matcher = sm.SwitchMatcher(
        setting_index.records, overrides=overrides, cb_alt_names=cb_alt, rules=rules
    )
    entry_by_id = {e.setting_id: e for e in matcher.index.entries}

    # Match per project: processed switches together, filtered ones apart.
    by_project: Dict[str, List[int]] = defaultdict(list)
    for i, row in enumerate(switch_rows):
        by_project[row.get("project", "")].append(i)
    refs = [to_ref(r) for r in switch_rows]
    results: List[Optional[sm.SwitchMatch]] = [None] * len(switch_rows)
    for project, idxs in by_project.items():
        for group in ([i for i in idxs if is_processed(switch_rows[i])],
                      [i for i in idxs if not is_processed(switch_rows[i])]):
            for i, res in zip(group, matcher.match_all([refs[i] for i in group])):
                results[i] = res

    # ---------------------------------------------------------------- switches
    switches = []
    legacy_changes = []
    claimed_by_project: Dict[str, Set[str]] = defaultdict(set)
    claimed_any: Set[str] = set()
    filtered_claims: Dict[str, List[str]] = defaultdict(list)
    for row, ref, res in zip(switch_rows, refs, results):
        processed = is_processed(row)
        legacy = legacy_ids(setting_index, row, cb_alt) if processed else []
        new = res.setting_ids if processed else []
        if processed and res.status == sm.MATCHED:
            claimed_by_project[ref.project].update(new)
            claimed_any.update(new)
        if not processed and res.status == sm.MATCHED:
            for sid in res.setting_ids:
                filtered_claims[sid].append(f"{ref.project} / {ref.name} ({filter_reason(row)})")

        change = ""
        if processed:
            if legacy == new:
                change = "same" if new else ""
            elif set(legacy) == set(new):
                change = "reordered"
            elif not legacy:
                change = "gained"
            elif not new:
                change = "lost"
            elif set(new) >= set(legacy):
                change = "extra records"
            else:
                change = "different records"
        rec = {
            "project": ref.project,
            "class": ref.cls,
            "container": ref.container,
            "switch_name": ref.name,
            "processed": int(processed),
            "filter_reason": "" if processed else filter_reason(row),
            "legacy_count": len(legacy) if processed else "",
            "status": res.status,
            "rule": res.rule,
            "ips_names": "; ".join(sorted({e.name.strip() for e in res.entries})),
            "ips_patterns": "; ".join(e.pattern for e in res.entries),
            "count": len(res.entries),
            "change_vs_legacy": change,
            "rejections": " | ".join(f"{j.pf_rule} '{j.pf_name}': {j.reason}" for j in res.rejections),
            "note": res.note,
            "override_problem": res.override_problem,
            "cb_alt_problem": res.cb_alt_problem,
            "relays": row.get("relays", ""),
            "fuses": row.get("fuses", ""),
            "aUsage": row.get("aUsage", ""),
            "grid": row.get("grid", ""),
            "study_case": row.get("study_case", ""),
            "full_name": ref.key,
            "setting_ids": "; ".join(res.setting_ids),
            "legacy_setting_ids": "; ".join(legacy),
        }
        switches.append(rec)
        if change in ("gained", "lost", "extra records", "different records",
                      "reordered"):
            lost = [entry_by_id[s] for s in legacy if s not in new and s in entry_by_id]
            added = [e for e in res.entries if e.setting_id not in legacy]
            legacy_changes.append({
                "project": ref.project, "class": ref.cls, "container": ref.container,
                "switch_name": ref.name, "change": change, "rule": res.rule,
                "records_added": "; ".join(_ips_label(e) for e in added),
                "records_removed": "; ".join(_ips_label(e) for e in lost),
                "why_removed": res.rejections[0].reason if lost and res.rejections else "",
                "relays": row.get("relays", ""),
            })

    # --------------------------------------------------------- relays at risk
    relays_at_risk = [
        s for s in switches
        if s["processed"] and s["status"] not in (sm.MATCHED, sm.REVIEWED_NONE)
        and (s["relays"] or s["fuses"])
    ]

    # ------------------------------------------------- filtered with IPS data
    filtered_with_ips = [
        {k: s[k] for k in ("project", "class", "container", "switch_name",
                           "filter_reason", "rule", "ips_names", "ips_patterns",
                           "relays", "study_case", "grid")}
        for s in switches if not s["processed"] and s["status"] == sm.MATCHED
    ]

    # ----------------------------------------------------------------- review
    review = []
    for s in switches:
        if not s["processed"]:
            continue
        if s["override_problem"]:
            review.append({"issue": "override target missing", **_short(s)})
        elif s["status"] == sm.AMBIGUOUS:
            review.append({"issue": s["status"], **_short(s)})
        elif s["status"] == sm.UNMATCHED and s["rejections"]:
            review.append({"issue": "candidate rejected", **_short(s)})

    # ------------------------------------------------------------ suggestions
    lookup = _Lookup(matcher.index.entries)
    off_rules = set(sm.ALL_RULES) - set(rules)
    off_matcher = sm.SwitchMatcher(
        setting_index.records, overrides=overrides, cb_alt_names=cb_alt,
        rules=set(sm.ALL_RULES),
    ) if off_rules else None
    suggestions = []
    for row, ref, res, s in zip(switch_rows, refs, results, switches):
        if not s["processed"] or res.status not in (sm.UNMATCHED, sm.AMBIGUOUS):
            continue
        slash = sm._DOWN_LINE_SLASH.match(ref.name.strip())
        protective = (
            res.status == sm.AMBIGUOUS
            or row.get("relays") or row.get("fuses")
            or (ref.cls == "ElmCoup" and row.get("aUsage") == "cbk"
                and ref.container_class == "ElmSubstat"
                and not ref.name.startswith("External Grid"))
            or (slash and re.sub(r"\d+$", "", slash.group(2).upper())
                in PROTECTIVE_SUFFIXES)
        )
        if protective:
            suggestions.extend(suggest(
                row, ref, res, lookup, claimed_by_project[ref.project],
                off_rules, off_matcher,
            ))
    suggested_for = defaultdict(list)
    for sg in suggestions:
        suggested_for[(sg["ips_name"].upper(), sg["ips_site"])].append(
            f"{sg['project']} / {sg['switch_name']}")

    # ---------------------------------------------------------- IPS unmatched
    sub_sites = {r.container for r in refs if r.cls == "ElmCoup"} | {
        r.substation for r in refs if r.substation}
    dl_sites: Set[str] = set()
    for r in refs:
        for _, name in sm._candidates(r, sm.ALL_RULES, {}):
            dl_sites.add(name.strip().upper())
    rejected_for: Dict[str, List[str]] = defaultdict(list)
    for s, res in zip(switches, results):
        if s["processed"]:
            for j in res.rejections:
                for sid in j.setting_ids:
                    rejected_for[sid].append(f"{s['project']} / {s['switch_name']}: {j.reason}")

    # Feeder bases (cable-box letter dropped) that some switch already took.
    matched_bases: Set[Tuple[str, str]] = set()
    for sid in claimed_any:
        e = entry_by_id.get(sid)
        if e is not None and e.site:
            for tok in [e.name.strip()] + sm._composite_keys(e.name):
                base = sm.coupler_base_name(tok.upper()) or tok.upper()
                matched_bases.add((base, e.site))

    ips_unmatched, ips_not_expected = [], []
    excluded = Counter()
    for e in matcher.index.entries:
        if e.category not in sm.IN_SCOPE_CATEGORIES:
            excluded[f"category '{e.category}'"] += 1
            continue
        if e.site is None:
            excluded["substation mapped to skip"] += 1
            continue
        if e.setting_id in claimed_any:
            continue
        rec = {
            "ips_name": e.name, "site": e.site, "category": e.category,
            "voltage": e.kv, "pattern": e.pattern, "element_type": "",
            "reason": "", "detail": "", "setting_id": e.setting_id,
            "date_setting": sm._get(e.record, "datesetting"),
        }
        if e.category == sm.CATEGORY_SUBSTATIONS:
            etype = classify_designation(sm._composite_keys(e.name)[0]
                                         if "+" in e.name else e.name.strip(), e.kv)
            rec["element_type"] = etype.value
            if (isinstance(e.kv, (int, float)) and e.kv >= 33) or etype in NOT_SWITCH_TYPES:
                rec["reason"] = "not modelled as a switch in distribution models"
                ips_not_expected.append(rec)
                continue
            present = e.site in sub_sites
        else:
            rec["element_type"] = "Down Line Device"
            present = e.site.upper() in dl_sites or e.name.strip().upper() in dl_sites
        key = (e.name.strip().upper(), e.site)
        if e.setting_id in filtered_claims:
            rec["reason"] = "only a switch IPStoPF skips matches it"
            rec["detail"] = "; ".join(filtered_claims[e.setting_id][:3])
        elif e.setting_id in rejected_for:
            rec["reason"] = "a candidate switch was rejected"
            rec["detail"] = "; ".join(rejected_for[e.setting_id][:3])
        elif key in suggested_for:
            rec["reason"] = "suggested override"
            rec["detail"] = "; ".join(suggested_for[key][:3])
        elif not present:
            rec["reason"] = "site not in any SEQ distribution model"
        elif e.category == sm.CATEGORY_SUBSTATIONS and any(
            ((sm.coupler_base_name(t.upper()) or t.upper()), e.site) in matched_bases
            for t in [e.name.strip()] + sm._composite_keys(e.name)
        ):
            rec["reason"] = "other cable box of a feeder that matched"
        else:
            rec["reason"] = "no switch with this name at the site"
        ips_unmatched.append(rec)

    # ------------------------------------------------------------ CB_ALT_NAME
    cb_alt_review = review_cb_alt_name(
        cb_alt_rows or [], switch_rows, refs, results, setting_index,
        overrides, cb_alt, rules,
    )
    for s, res in zip(switches, results):
        if s["processed"] and res.cb_alt_problem:
            review.append({"issue": "CB_ALT_NAME entry finds nothing", **_short(s)})

    # ------------------------------------------------------------- per project
    projects = []
    for project, idxs in sorted(by_project.items()):
        ss = [switches[i] for i in idxs]
        proc = [s for s in ss if s["processed"]]
        case = Counter(s["study_case"] for s in ss).most_common(1)
        projects.append({
            "project": project,
            "study_case": case[0][0] if case else "",
            "switches": len(ss),
            "processed": len(proc),
            "legacy_matched": sum(1 for s in proc if s["legacy_count"]),
            "new_matched": sum(1 for s in proc if s["status"] == sm.MATCHED),
            "gained": sum(1 for s in proc if s["change_vs_legacy"] == "gained"),
            "lost": sum(1 for s in proc if s["change_vs_legacy"] == "lost"),
            "extra_records": sum(1 for s in proc if s["change_vs_legacy"] == "extra records"),
            "reordered": sum(1 for s in proc if s["change_vs_legacy"] == "reordered"),
            "ambiguous": sum(1 for s in proc if s["status"] == sm.AMBIGUOUS),
            "relays_at_risk": sum(1 for s in relays_at_risk if s["project"] == project),
            "filtered_with_ips": sum(1 for s in filtered_with_ips if s["project"] == project),
        })

    return {
        "switches": switches,
        "relays_at_risk": relays_at_risk,
        "ips_unmatched": ips_unmatched,
        "ips_not_expected": ips_not_expected,
        "filtered_with_ips": filtered_with_ips,
        "review": review,
        "legacy_changes": legacy_changes,
        "suggested_overrides": suggestions,
        "cb_alt_name": cb_alt_review,
        "projects": projects,
        "_excluded": excluded,
        "_ips_records": len(matcher.index.entries),
        "_ips_raw": len(ips_rows),
        "_rules": sorted(rules),
        "_overrides": len(overrides),
    }


def review_cb_alt_name(
    cb_alt_rows: List[Dict[str, str]],
    switch_rows: List[Dict[str, str]],
    refs: List[sm.SwitchRef],
    results: List[sm.SwitchMatch],
    setting_index: Any,
    overrides: Sequence[sm.Override],
    cb_alt: Dict,
    rules: Set[str],
) -> List[Dict[str, str]]:
    """
    What each CB_ALT_NAME.csv row does today, and what to do with it.

    The file is keyed on (SUBSTATION, CB_NAME) only - PROJECT and GRID are
    not used - so a row applies wherever that coupler exists.
    """
    if not cb_alt_rows:
        return []
    bare = sm.SwitchMatcher(
        setting_index.records, overrides=overrides, cb_alt_names={}, rules=rules)
    entry_by_id = {e.setting_id: e for e in bare.index.entries}

    def lab(ids):
        return "; ".join(sorted({f"{entry_by_id[i].name.strip()} @ {entry_by_id[i].site}"
                                 for i in ids if i in entry_by_id})) or "-"

    by_key: Dict[Tuple[str, str], List[int]] = defaultdict(list)
    for i, row in enumerate(switch_rows):
        by_key[(row.get("fold_id_name", ""), row.get("loc_name", ""))].append(i)

    # The same project matched without CB_ALT_NAME, for comparison.
    keyed = {i for r in cb_alt_rows for i in by_key.get((r.get("SUBSTATION"), r.get("CB_NAME")), [])}
    projects = {switch_rows[i].get("project", "") for i in keyed}
    without: Dict[int, sm.SwitchMatch] = {}
    for project in projects:
        idxs = [i for i, r in enumerate(switch_rows)
                if r.get("project") == project and is_processed(r)]
        for i, res in zip(idxs, bare.match_all([refs[i] for i in idxs])):
            without[i] = res

    seen: Set[Tuple[str, str]] = set()
    out = []
    for row in cb_alt_rows:
        key = (row.get("SUBSTATION", ""), row.get("CB_NAME", ""))
        new_name = row.get("NEW_NAME", "").strip()
        is_note = new_name.lower() in CB_ALT_NOTES
        rec = {
            "line": row.get("line", ""), "PROJECT": row.get("PROJECT", ""),
            "SUBSTATION": key[0], "CB_NAME": key[1], "NEW_NAME": new_name,
            "row_type": "note" if is_note else "rename",
            "found_in": "", "legacy_with_row": "", "legacy_without_row": "",
            "new_result": "", "new_without_row": "", "effect": "", "action": "",
        }
        out.append(rec)
        # Is this the row ex_settings actually uses for the key?
        effective = cb_alt.get(key) == row.get("NEW_NAME") and key not in seen
        duplicate = key in cb_alt and not effective and not (
            new_name.lower() in CB_ALT_NOTES - {"not required", "decommisioned",
                                                "decommissioned"})
        if effective:
            seen.add(key)
        hits = by_key.get(key, [])
        rec["found_in"] = "; ".join(sorted({switch_rows[i].get("project", "") for i in hits}))
        if not hits:
            rec["effect"] = "switch not in any SEQ model"
            rec["action"] = "optional: delete (switch not in any model)"
            continue
        processed = [i for i in hits if is_processed(switch_rows[i])]
        if not processed:
            rec["effect"] = "switch skipped by IPStoPF: " + filter_reason(switch_rows[hits[0]])
            rec["action"] = "no effect while the switch is skipped"
            continue
        i = processed[0]
        res, bare_res = results[i], without.get(i)
        lw = legacy_ids(setting_index, switch_rows[i], cb_alt)
        lwo = legacy_ids(setting_index, switch_rows[i], {})
        rec["legacy_with_row"] = lab(lw)
        rec["legacy_without_row"] = lab(lwo)
        rec["new_result"] = f"{res.status} {res.rule} {lab(res.setting_ids)}".strip()
        if bare_res is not None:
            rec["new_without_row"] = (
                f"{bare_res.status} {bare_res.rule} {lab(bare_res.setting_ids)}".strip())

        if duplicate:
            rec["effect"] = ("ignored: an earlier row has the same SUBSTATION and "
                             "CB_NAME (PROJECT is not used)")
            same = (cb_alt.get(key) or "").split("_")[0] == row.get("NEW_NAME", "").split("_")[0]
            rec["action"] = ("optional: delete (duplicate)" if same
                             else "resolve: conflicts with the earlier row")
        elif effective and set(lw) < set(lwo):
            rec["effect"] = ("breaks a match: without this row the switch matches by "
                             "its own name; with it, legacy finds nothing and deletes "
                             "the relays")
            rec["action"] = "DELETE" if not is_note else "change NEW_NAME to Not Needed"
        elif not is_note:
            if res.cb_alt_problem:
                rec["effect"] = "NEW_NAME finds no IPS record at this substation"
                rec["action"] = "fix NEW_NAME or delete"
            elif bare_res is not None and bare_res.setting_ids == res.setting_ids:
                rec["effect"] = "new rules give the same records without it"
                rec["action"] = "optional: delete once on new mode (harmless to keep)"
            else:
                rec["effect"] = "needed: maps the switch to " + lab(res.setting_ids)
                rec["action"] = "keep"
        else:
            note = new_name.lower()
            if res.status == sm.MATCHED and lw:
                rec["effect"] = "switch matches today by its own name; the note has no effect"
                rec["action"] = ("update note: IPS has a setting" if note == "no active setting"
                                 else "review note (no effect on matching)")
            elif res.status == sm.MATCHED:
                rec["effect"] = f"new rule '{res.rule}' now matches it"
                rec["action"] = (
                    "resolved by the new matcher; replace the note" if note == "unknown"
                    else "review: add a NONE override if it should stay unmatched")
            else:
                rec["effect"] = "unmatched (as in legacy)"
                rec["action"] = "no change" if note != "unknown" else "still unknown"
    return out


def _short(s: Dict[str, Any]) -> Dict[str, Any]:
    return {k: s[k] for k in ("project", "class", "container", "switch_name",
                              "status", "rule", "rejections", "note",
                              "override_problem", "relays")}


# =============================================================================
# Output
# =============================================================================

SHEETS = [
    ("relays_at_risk", "Relays at risk"),
    ("suggested_overrides", "Suggested overrides"),
    ("cb_alt_name", "CB_ALT_NAME"),
    ("review", "Review"),
    ("legacy_changes", "Changes vs legacy"),
    ("filtered_with_ips", "Filtered with IPS"),
    ("ips_unmatched", "IPS unmatched"),
    ("ips_not_expected", "IPS not expected"),
    ("projects", "Projects"),
    ("switches", "All switches"),
]


def _write_csv(path: Path, rows: List[Dict[str, Any]]) -> None:
    fields: List[str] = []
    for r in rows:
        for k in r:
            if k not in fields:
                fields.append(k)
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=fields or ["(none)"], quoting=csv.QUOTE_ALL)
        w.writeheader()
        w.writerows(rows)


def _write_xlsx(path: Path, report: Dict[str, Any]) -> bool:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
        from openpyxl.utils import get_column_letter
    except ImportError:
        return False
    wb = Workbook()
    wb.remove(wb.active)
    for key, title in SHEETS:
        rows = report[key]
        ws = wb.create_sheet(title)
        fields: List[str] = []
        for r in rows:
            for k in r:
                if k not in fields:
                    fields.append(k)
        ws.append(fields or ["(none)"])
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in rows:
            ws.append([("" if r.get(k) is None else str(r.get(k)))[:32000] for k in fields])
        ws.freeze_panes = "A2"
        if fields:
            ws.auto_filter.ref = ws.dimensions
        for i, k in enumerate(fields, 1):
            width = max([len(k)] + [len(str(r.get(k, ""))) for r in rows[:500]])
            ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 8), 60)
    wb.save(path)
    return True


def write_summary(path: Path, report: Dict[str, Any], sources: Dict[str, str]) -> str:
    sw = report["switches"]
    proc = [s for s in sw if s["processed"]]
    status = Counter(s["status"] for s in proc)
    rules = Counter(s["rule"] for s in proc if s["status"] == sm.MATCHED)
    change = Counter(s["change_vs_legacy"] for s in proc if s["change_vs_legacy"])
    unmatched = Counter(r["reason"] for r in report["ips_unmatched"])
    sugg = Counter(s["confidence"] for s in report["suggested_overrides"])
    odd_cases = [p for p in report["projects"]
                 if p["study_case"] and p["study_case"] != ALL_GRIDS_STUDY_CASE]

    lines = [
        "# Switch match audit",
        "",
        f"Run {datetime.now():%Y-%m-%d %H:%M}.",
        "",
        f"* Switches: {sources['switches']}",
        f"* IPS: {sources['ips']} ({report['_ips_raw']} rows, "
        f"{report['_ips_records']} after archived/excluded filtering)",
        f"* Overrides: {sources['overrides']} ({report['_overrides']} rows)",
        f"* Rules: {', '.join(report['_rules'])}",
        "",
        "## Switches",
        "",
        f"{len(sw)} switches in {len(report['projects'])} projects; "
        f"{len(proc)} pass IPStoPF's filters and are matched.",
        "",
        "| Result | Switches |",
        "|---|---:|",
    ]
    lines += [f"| {k} | {v} |" for k, v in status.most_common()]
    lines += ["", "Matched by rule:", "", "| Rule | Switches |", "|---|---:|"]
    lines += [f"| {k} | {v} |" for k, v in rules.most_common()]
    lines += ["", "Compared with the legacy lookup:", "", "| Change | Switches |", "|---|---:|"]
    lines += [f"| {k} | {v} |" for k, v in change.most_common()]
    lines += [
        "",
        "## To look at",
        "",
        f"* **Relays at risk**: {len(report['relays_at_risk'])} processed switches "
        f"have relays or fuses in their cubicle but no IPS match. IPStoPF deletes "
        f"those relays.",
        f"* **Suggested overrides**: {len(report['suggested_overrides'])} "
        f"({', '.join(f'{v} {k}' for k, v in sugg.most_common()) or 'none'}). "
        f"Review, then copy the rows you agree with into switch_match_overrides.csv.",
        f"* **Review**: {len(report['review'])} ambiguous or rejected matches.",
        f"* **Filtered with IPS**: {len(report['filtered_with_ips'])} switches "
        f"IPStoPF skips although IPS has in-service settings for them.",
        f"* **IPS unmatched**: {len(report['ips_unmatched'])} in-scope IPS records "
        f"no switch takes ({len(report['ips_not_expected'])} more are substation "
        f"elements the distribution models don't carry as switches).",
    ]
    lines += [f"    * {k}: {v}" for k, v in unmatched.most_common()]
    if odd_cases:
        lines += ["", "## Study case", ""]
        for p in odd_cases:
            lines.append(
                f"* {p['project']} opened in '{p['study_case']}', not "
                f"'{ALL_GRIDS_STUDY_CASE}'. IPStoPF skips every switch that is "
                f"not energised in the active study case "
                f"({p['filtered_with_ips']} of them have IPS settings)."
            )
    text = "\n".join(lines) + "\n"
    path.write_text(text, encoding="utf-8")
    return text


def main(argv: Optional[Sequence[str]] = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--switches", required=True,
                    help="export_pf_switches.py CSV, or a folder of switch_match_*.csv")
    ap.add_argument("--ips", help="Report-Cache-ProtectionSettingIDs-EX.csv")
    ap.add_argument("--overrides", help="switch_match_overrides.csv")
    ap.add_argument("--out", help="report folder (default: next to --switches)")
    ap.add_argument("--enable", nargs="*", default=[], help="extra rules")
    ap.add_argument("--disable", nargs="*", default=[], help="rules to switch off")
    ap.add_argument("--cb-alt-name", help="CB_ALT_NAME.csv (default: the configured one)")
    ap.add_argument("--no-cb-alt-name", action="store_true",
                    help="ignore CB_ALT_NAME.csv")
    args = ap.parse_args(argv)

    switch_rows, files = load_switch_rows(args.switches)
    ips_rows, ips_source = load_ips_rows(args.ips)
    ov_path = args.overrides or str(smc.get_overrides_file())
    overrides, problems = sm.load_overrides(ov_path)
    for p in problems:
        print(f"overrides {p}; row ignored")
    cb_alt_rows = [] if args.no_cb_alt_name else load_cb_alt_rows(args.cb_alt_name)
    cb_alt = cb_alt_index(cb_alt_rows)
    rules = (set(sm.DEFAULT_RULES) | set(smc.ENABLE_RULES) | set(args.enable)) - (
        set(smc.DISABLE_RULES) | set(args.disable))
    unknown = rules - sm.ALL_RULES
    if unknown:
        raise SystemExit(f"Unknown rule(s): {sorted(unknown)}; known: {sorted(sm.ALL_RULES)}")

    from update_powerfactory import mapping_file
    mapping_file._load_type_mapping()
    if not mapping_file._excluded_patterns_cache:
        print("WARNING: type_mapping.csv gave no excluded patterns; check it was found")

    report = run_audit(switch_rows, ips_rows, overrides, cb_alt, rules, cb_alt_rows)

    base = Path(args.out) if args.out else (
        Path(args.switches) if Path(args.switches).is_dir() else Path(args.switches).parent)
    out = base / f"switch_match_audit_{datetime.now():%Y%m%d-%H%M%S}"
    out.mkdir(parents=True, exist_ok=True)
    for key, _ in SHEETS:
        _write_csv(out / f"{key}.csv", report[key])
    xlsx = _write_xlsx(out / "switch_match_audit.xlsx", report)
    text = write_summary(out / "summary.md", report, {
        "switches": f"{len(files)} file(s), {len(switch_rows)} rows"
                    + (f" ({files[0]})" if len(files) == 1 else f" ({args.switches})"),
        "ips": ips_source,
        "overrides": f"{ov_path}" + ("" if os.path.exists(ov_path) else " (not found)"),
    })
    print(text)
    print(f"Report: {out}" + ("" if xlsx else " (openpyxl missing: CSVs only)"))
    return 0


if __name__ == "__main__":
    sys.exit(main())
